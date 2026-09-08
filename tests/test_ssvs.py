import os
import unittest
from app import create_app
from app.models import db
from app.models.user import Teacher
from app.models.form import Form, FormField, FieldOption
from app.models.submission import Submission
from app.models.scoring import ScoreFormula, FormulaRule
from app.extractors.base import BaseExtractor
from app.extractors.mock_simulator import MockSimulator
from app.scoring_engine.formula_evaluator import FormulaEvaluator
from app.scoring_engine.normalizer import Normalizer
from app.export.excel_exporter import ExcelExporter
import openpyxl

class SSVSTestCase(unittest.TestCase):
    def setUp(self):
        self.app = create_app('testing')
        self.app_context = self.app.app_context()
        self.app_context.push()
        self.client = self.app.test_client()
        db.create_all()

        # Seed test teacher
        self.teacher = Teacher(
            email='test@ssvs.edu',
            full_name='Dr. Test Evaluator',
            college_name='Testing Institute of Technology',
            department='Computer Science'
        )
        self.teacher.set_password('SecretPass123!')
        db.session.add(self.teacher)
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()

    def test_teacher_password_hashing(self):
        """Verify password hashing and verification."""
        self.assertTrue(self.teacher.check_password('SecretPass123!'))
        self.assertFalse(self.teacher.check_password('WrongPassword'))

    def test_url_extractor_parsing(self):
        """Verify regex parsing of platform profile URLs."""
        # LeetCode
        self.assertEqual(BaseExtractor.extract_username('https://leetcode.com/u/neal_wu/', 'leetcode'), 'neal_wu')
        self.assertEqual(BaseExtractor.extract_username('leetcode.com/john_doe', 'leetcode'), 'john_doe')
        self.assertEqual(BaseExtractor.extract_username('plain_user', 'leetcode'), 'plain_user')
        
        # Codeforces
        self.assertEqual(BaseExtractor.extract_username('https://codeforces.com/profile/tourist', 'codeforces'), 'tourist')
        
        # GitHub
        self.assertEqual(BaseExtractor.extract_username('https://github.com/octocat/', 'github'), 'octocat')
        
        # CodeChef
        self.assertEqual(BaseExtractor.extract_username('https://www.codechef.com/users/chef_pro', 'codechef'), 'chef_pro')

    def test_mock_simulator_determinism(self):
        """Verify simulator produces consistent realistic metrics for identical usernames."""
        s1 = MockSimulator.simulate_leetcode('neal_wu')
        s2 = MockSimulator.simulate_leetcode('neal_wu')
        self.assertEqual(s1.total_solved, s2.total_solved)
        self.assertEqual(s1.contest_rating, s2.contest_rating)
        self.assertGreater(s1.total_solved, 0)
        self.assertGreater(s1.contest_rating, 1000)

    def test_scoring_formula_and_penalties(self):
        """Verify weighted formula scoring, bonus rules, and backlog penalties."""
        form = Form(
            teacher_id=self.teacher.id,
            title='Test Assessment',
            slug='test-assessment-slug',
            is_published=True
        )
        db.session.add(form)
        db.session.flush()

        formula = ScoreFormula(
            form_id=form.id,
            name='Test Formula',
            max_total_marks=100.0,
            is_active=True
        )
        db.session.add(formula)
        db.session.flush()

        # CGPA: 10 CGPA = 30 marks
        db.session.add(FormulaRule(
            formula_id=formula.id,
            metric_key='cgpa',
            display_label='Academic CGPA',
            weight_percentage=30.0,
            max_marks=30.0,
            multiplier=3.0,
            category='academic'
        ))

        # Backlogs: -5 marks per backlog
        db.session.add(FormulaRule(
            formula_id=formula.id,
            metric_key='backlogs_penalty',
            display_label='Backlogs Penalty',
            weight_percentage=0.0,
            max_marks=0.0,
            penalty_per_unit=5.0,
            category='penalty'
        ))
        db.session.commit()

        # Create submission with CGPA 9.0 and 2 backlogs:
        # Earned = 9.0 * 3.0 = 27.0
        # Deduction = 2 * 5.0 = 10.0
        # Expected total = 17.0
        sub = Submission(
            form_id=form.id,
            student_name='Test Student',
            roll_number='2026TEST01',
            email='student@test.edu',
            cgpa=9.0,
            backlogs=2,
            status='completed'
        )
        db.session.add(sub)
        db.session.commit()

        calc = FormulaEvaluator.evaluate_submission(sub, formula)
        self.assertEqual(calc.academic_score, 27.0)
        self.assertEqual(calc.penalty_deductions, 10.0)
        self.assertEqual(calc.total_score, 17.0)

    def test_normalizer(self):
        """Verify min-max and percentile calculations."""
        scaled = Normalizer.min_max(75, 50, 100, 0, 100)
        self.assertEqual(scaled, 50.0)

        scores = [10.0, 20.0, 30.0, 40.0, 50.0]
        percentiles = Normalizer.calculate_percentiles(scores)
        self.assertEqual(len(percentiles), 5)
        self.assertGreater(percentiles[-1], percentiles[0])

    def test_excel_exporter(self):
        """Verify Excel workbook creates all 4 required sheets with valid structure."""
        form = Form(
            teacher_id=self.teacher.id,
            title='Campus Placement Test Drive',
            slug='test-drive-slug',
            is_published=True
        )
        db.session.add(form)
        db.session.flush()

        sub = Submission(
            form_id=form.id,
            student_name='Sample Candidate',
            roll_number='ROLL001',
            email='sample@college.edu',
            cgpa=8.5,
            backlogs=0
        )
        db.session.add(sub)
        db.session.commit()

        FormulaEvaluator.update_form_ranks_and_leaderboard(form.id)

        stream = ExcelExporter.generate_form_report(form)
        self.assertIsNotNone(stream)

        # Inspect workbook
        wb = openpyxl.load_workbook(stream)
        sheet_names = wb.sheetnames
        self.assertIn("Leaderboard & Ranks", sheet_names)
        self.assertIn("Raw Submissions", sheet_names)
        self.assertIn("Coding Platform Stats", sheet_names)
        self.assertIn("Executive Analytics", sheet_names)

    def test_public_pages_accessible_without_login(self):
        """Verify student zero-login accessibility for public endpoints."""
        resp = self.client.get('/leaderboard')
        self.assertEqual(resp.status_code, 200)

        form = Form(
            teacher_id=self.teacher.id,
            title='Open Form',
            slug='open-form-slug',
            is_published=True
        )
        db.session.add(form)
        db.session.commit()

        resp = self.client.get(f'/f/{form.slug}')
        self.assertEqual(resp.status_code, 200)

    def test_new_platform_extractors(self):
        """Verify URL parsing and mock simulation for AtCoder, InterviewBit, and Kaggle."""
        # URL Parsing
        self.assertEqual(BaseExtractor.extract_username('https://atcoder.jp/users/chokudai', 'atcoder'), 'chokudai')
        self.assertEqual(BaseExtractor.extract_username('https://www.interviewbit.com/profile/ib_coder_99', 'interviewbit'), 'ib_coder_99')
        self.assertEqual(BaseExtractor.extract_username('https://www.kaggle.com/data_wizard', 'kaggle'), 'data_wizard')

        # Mock simulation
        atc = MockSimulator.simulate_atcoder('chokudai')
        self.assertGreater(atc.current_rating, 0)
        self.assertIn(atc.tier_color, ['White', 'Brown', 'Green', 'Cyan', 'Blue', 'Yellow'])

        ib = MockSimulator.simulate_interviewbit('ib_coder_99')
        self.assertGreater(ib.score, 0)

        kg = MockSimulator.simulate_kaggle('data_wizard')
        self.assertGreater(kg.total_medals, 0)
        self.assertIn(kg.tier, ['Novice', 'Contributor', 'Expert', 'Master', 'Grandmaster'])

    def test_custom_excel_export_rows_and_columns(self):
        """Verify custom Excel export contains ONLY chosen columns and rows."""
        form = Form(
            teacher_id=self.teacher.id,
            title='Custom Export Test Drive',
            slug='custom-export-slug',
            is_published=True
        )
        db.session.add(form)
        db.session.flush()

        s1 = Submission(form_id=form.id, student_name='Alpha One', roll_number='R001', email='alpha@test.edu', cgpa=9.2)
        s2 = Submission(form_id=form.id, student_name='Beta Two', roll_number='R002', email='beta@test.edu', cgpa=8.4)
        s3 = Submission(form_id=form.id, student_name='Gamma Three', roll_number='R003', email='gamma@test.edu', cgpa=7.8)
        db.session.add_all([s1, s2, s3])
        db.session.commit()

        FormulaEvaluator.update_form_ranks_and_leaderboard(form.id)

        # Export only s1 and s2, with columns ['student_name', 'roll_number', 'cgpa']
        stream = ExcelExporter.generate_custom_export(
            form=form,
            submission_ids=[s1.id, s2.id],
            column_keys=['student_name', 'roll_number', 'cgpa']
        )
        self.assertIsNotNone(stream)

        wb = openpyxl.load_workbook(stream)
        ws = wb.active
        self.assertEqual(ws.title, "Custom Export")
        
        # Row 3 should be column headers: Student Name, Roll Number, CGPA
        header_vals = [ws.cell(row=3, column=c).value for c in range(1, 4)]
        self.assertEqual(header_vals, ["Student Name", "Roll Number", "CGPA"])
        
        # Row 4 and Row 5 are the 2 data rows
        # Check that only 2 data rows exist
        self.assertEqual(ws.max_row, 5)

    def test_teacher_edit_submission_and_manual_adjustment(self):
        """Verify teacher editing student data and manual score adjustments."""
        form = Form(
            teacher_id=self.teacher.id,
            title='Placement Assessment',
            slug='placement-assessment',
            is_published=True
        )
        db.session.add(form)
        db.session.flush()

        formula = ScoreFormula(form_id=form.id, name='Simple Formula', is_active=True, max_total_marks=100.0)
        db.session.add(formula)
        db.session.flush()
        db.session.add(FormulaRule(formula_id=formula.id, metric_key='cgpa', display_label='CGPA', weight_percentage=100.0, max_marks=100.0, multiplier=10.0, category='academic'))
        db.session.commit()

        sub = Submission(form_id=form.id, student_name='Candidate A', roll_number='R100', email='cand@test.edu', cgpa=8.0)
        db.session.add(sub)
        db.session.commit()

        # Initial calculation: 8.0 * 10 = 80.0
        calc = FormulaEvaluator.evaluate_submission(sub, formula)
        self.assertEqual(calc.total_score, 80.0)

        # Teacher modifies student data and adds +5 manual score adjustment
        sub.cgpa = 8.5
        sub.manual_score_adjustment = 5.0
        sub.verification_status = 'shortlisted'
        sub.teacher_remarks = 'Excellent interview performance'
        db.session.commit()

        # Recalculate: (8.5 * 10) + 5.0 = 90.0
        calc_updated = FormulaEvaluator.evaluate_submission(sub, formula)
        self.assertEqual(calc_updated.total_score, 90.0)
        self.assertEqual(sub.verification_status, 'shortlisted')

    def test_scorecard_certificate_endpoint(self):
        """Verify student official single-page scorecard/certificate rendering."""
        form = Form(
            teacher_id=self.teacher.id,
            title='Scorecard Certification Form',
            slug='scorecard-cert-slug',
            is_published=True
        )
        db.session.add(form)
        db.session.flush()

        formula = ScoreFormula(form_id=form.id, name='Placement Metric', is_active=True, max_total_marks=100.0)
        db.session.add(formula)
        db.session.flush()

        sub = Submission(
            form_id=form.id,
            student_name='Dev Sharma',
            roll_number='CS2026-99',
            email='dev@college.edu',
            department='CSE',
            cgpa=9.5,
            backlogs=0,
            status='completed'
        )
        db.session.add(sub)
        db.session.commit()

        calc = FormulaEvaluator.evaluate_submission(sub, formula)

        resp = self.client.get(f'/submission/{sub.uuid}/scorecard')
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        self.assertIn('Dev Sharma', html)
        self.assertIn('CS2026-99', html)
        self.assertIn('Print Scorecard', html)
        self.assertIn('data:image/png;base64,', html)
        self.assertIn('Official Performance Scorecard', html)

    def test_email_otp_lifecycle(self):
        """Test OTP generation, SHA-256 hashing, attempt tracking, and verification."""
        from app.models.otp import EmailOTP
        
        email = "pk0403564@gmail.com"
        record, code = EmailOTP.create_otp(email, purpose='registration', validity_minutes=10)
        
        # Verify 6-digit numeric format
        self.assertEqual(len(code), 6)
        self.assertTrue(code.isdigit())
        
        # Verify hash match
        self.assertEqual(record.otp_hash, EmailOTP.hash_code(code))
        self.assertFalse(record.is_used)
        
        # Test wrong OTP increments attempts
        valid, msg = EmailOTP.verify_otp(email, 'registration', '000000' if code != '000000' else '111111')
        self.assertFalse(valid)
        self.assertIn('attempt', msg)
        
        # Refresh record
        record = db.session.get(EmailOTP, record.id)
        self.assertEqual(record.attempts, 1)
        
        # Test correct OTP verifies
        valid, msg = EmailOTP.verify_otp(email, 'registration', code)
        self.assertTrue(valid)
        
        # Refresh record
        record = db.session.get(EmailOTP, record.id)
        self.assertTrue(record.is_used)
        
        # Cannot reuse verified OTP
        valid_again, _ = EmailOTP.verify_otp(email, 'registration', code)
        self.assertFalse(valid_again)

    def test_otp_cooldown_rate_limiting(self):
        """Test that requesting OTPs respects the 60-second cooldown."""
        from app.models.otp import EmailOTP
        
        email = "pk0403564@gmail.com"
        EmailOTP.create_otp(email, purpose='registration', validity_minutes=10)
        
        can_resend, remaining = EmailOTP.can_resend(email, 'registration', cooldown_seconds=60)
        self.assertFalse(can_resend)
        self.assertGreater(remaining, 0)
        self.assertLessEqual(remaining, 60)

    def test_auth_verify_email_and_password_reset(self):
        """Test complete auth verification flow with EmailOTP."""
        from app.models.otp import EmailOTP
        
        # Register a new unverified user
        new_teacher = Teacher(
            email='new_prof@ssvs.edu',
            full_name='Prof. New Teacher',
            college_name='IIT Delhi',
            department='CSE',
            email_verified=False
        )
        new_teacher.set_password('InitialPass123!')
        db.session.add(new_teacher)
        db.session.commit()
        
        # Generate OTP
        _, code = EmailOTP.create_otp(new_teacher.email, purpose='registration')
        
        # Verify via route
        with self.client.session_transaction() as sess:
            sess['pending_verification_email'] = new_teacher.email
            sess['pending_verification_purpose'] = 'registration'
            
        resp = self.client.post(
            f'/auth/verify-email?email={new_teacher.email}&purpose=registration',
            data={'otp_code': code},
            follow_redirects=True
        )
        self.assertEqual(resp.status_code, 200)
        
        # Verify user is now verified
        updated = Teacher.query.filter_by(email=new_teacher.email).first()
        self.assertTrue(updated.email_verified)
        
        # Log out before testing password reset
        self.client.get('/auth/logout')
        
        # Test password reset OTP
        _, reset_code = EmailOTP.create_otp(new_teacher.email, purpose='password_reset')
        resp = self.client.post(
            f'/auth/reset-password?email={new_teacher.email}',
            data={
                'email': new_teacher.email,
                'otp_code': reset_code,
                'password': 'BrandNewPassword2026!',
                'confirm_password': 'BrandNewPassword2026!'
            },
            follow_redirects=True
        )
        self.assertEqual(resp.status_code, 200)
        
        # Verify password changed
        updated = Teacher.query.filter_by(email=new_teacher.email).first()
        self.assertTrue(updated.check_password('BrandNewPassword2026!'))

    def test_teacher_login_otp_flow(self):
        """Test teacher login with 2-step SMTP OTP verification."""
        from app.models.otp import EmailOTP

        # 1. Login with password -> triggers OTP send and redirect to verify-otp
        resp = self.client.post('/auth/login', data={
            'email': self.teacher.email,
            'password': 'SecretPass123!'
        }, follow_redirects=False)
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/auth/verify', resp.headers['Location'])

        # Verify an active login OTP was generated
        latest_otp = EmailOTP.query.filter_by(email=self.teacher.email, purpose='login', is_used=False).first()
        self.assertIsNotNone(latest_otp)

        # 2. Complete verification with OTP code
        _, valid_code = EmailOTP.create_otp(self.teacher.email, purpose='login')
        resp = self.client.post(
            f'/auth/verify-otp?email={self.teacher.email}&purpose=login',
            data={'otp_code': valid_code},
            follow_redirects=True
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn('Verification successful!', resp.get_data(as_text=True))

    def test_delete_account_and_cascade_all_data(self):
        """Test that deleting a teacher account cascades and deletes all related forms, submissions, scores, logs, and OTPs."""
        from app.models.form import Form, FormField
        from app.models.submission import Submission, SubmissionValue
        from app.models.scoring import ScoreFormula, FormulaRule, CalculatedScore, LeaderboardEntry
        from app.models.audit import ActivityLog
        from app.models.otp import EmailOTP

        # 1. Create a dedicated teacher to be deleted
        teacher = Teacher(
            email='delete_me@ssvs.edu',
            full_name='Prof. To Be Deleted',
            college_name='Testing University',
            department='CSE'
        )
        teacher.set_password('DeletePass123!')
        db.session.add(teacher)
        db.session.commit()
        teacher_id = teacher.id

        # 2. Create form, field, submission, formula, scores, logs, and OTP
        form = Form(teacher_id=teacher_id, title='Assessment To Delete', slug='delete-slug-test')
        db.session.add(form)
        db.session.commit()
        form_id = form.id

        field = FormField(form_id=form_id, field_key='github_username', label='GitHub', field_type='github', display_order=1)
        db.session.add(field)

        formula = ScoreFormula(form_id=form_id, name='Formula To Delete', is_active=True)
        db.session.add(formula)
        db.session.commit()

        sub = Submission(form_id=form_id, student_name='Student To Delete', roll_number='DEL-01', email='del@student.edu')
        db.session.add(sub)
        db.session.commit()

        sub_val = SubmissionValue(submission_id=sub.id, field_id=field.id, value_text='torvalds')
        db.session.add(sub_val)

        calc = CalculatedScore(submission_id=sub.id, formula_id=formula.id, total_score=85.0)
        db.session.add(calc)

        lb = LeaderboardEntry(form_id=form_id, submission_id=sub.id, rank=1, student_name='Student To Delete', roll_number='DEL-01', total_score=85.0)
        db.session.add(lb)

        log = ActivityLog(teacher_id=teacher_id, action='TEST_LOG', description='Test audit')
        db.session.add(log)

        EmailOTP.create_otp('delete_me@ssvs.edu', purpose='login')
        db.session.commit()

        # 3. Authenticate as this teacher
        with self.client.session_transaction() as sess:
            sess['_user_id'] = str(teacher_id)
            sess['_fresh'] = True

        # 4. Wrong password should fail deletion
        resp = self.client.post('/auth/delete-account', data={
            'password': 'WrongPassword!',
            'confirm_phrase': 'DELETE'
        }, follow_redirects=True)
        self.assertIn('Incorrect password', resp.get_data(as_text=True))
        self.assertIsNotNone(db.session.get(Teacher, teacher_id))

        # 5. Wrong confirmation phrase should fail deletion
        resp = self.client.post('/auth/delete-account', data={
            'password': 'DeletePass123!',
            'confirm_phrase': 'CANCEL'
        }, follow_redirects=True)
        self.assertIn('You must type DELETE', resp.get_data(as_text=True))
        self.assertIsNotNone(db.session.get(Teacher, teacher_id))

        # 6. Correct password and confirmation phrase should delete account and cascade
        resp = self.client.post('/auth/delete-account', data={
            'password': 'DeletePass123!',
            'confirm_phrase': 'DELETE'
        }, follow_redirects=True)
        self.assertEqual(resp.status_code, 200)
        self.assertIn('permanently deleted', resp.get_data(as_text=True))

        # 7. Verify all related data is purged
        self.assertIsNone(db.session.get(Teacher, teacher_id))
        self.assertEqual(Form.query.filter_by(teacher_id=teacher_id).count(), 0)
        self.assertEqual(FormField.query.filter_by(form_id=form_id).count(), 0)
        self.assertEqual(Submission.query.filter_by(form_id=form_id).count(), 0)
        self.assertEqual(SubmissionValue.query.filter_by(submission_id=sub.id).count(), 0)
        self.assertEqual(ScoreFormula.query.filter_by(form_id=form_id).count(), 0)
        self.assertEqual(CalculatedScore.query.filter_by(submission_id=sub.id).count(), 0)
        self.assertEqual(LeaderboardEntry.query.filter_by(form_id=form_id).count(), 0)
        self.assertEqual(ActivityLog.query.filter_by(teacher_id=teacher_id).count(), 0)
        self.assertEqual(EmailOTP.query.filter_by(email='delete_me@ssvs.edu').count(), 0)

    def test_smtp_email_service_and_mock_dispatch(self):
        """Test SMTP email sending, recipient validation, and SMTP mock transport."""
        from unittest.mock import patch, MagicMock
        from app.utils.email_service import send_smtp_email, send_otp_code_email, get_smtp_config

        # 1. Invalid recipient email rejection
        sent, err = send_smtp_email("", "Subject", "<p>Body</p>")
        self.assertFalse(sent)
        self.assertIn("Invalid recipient", err)

        # 2. Testing mode mock dispatch (current app has TESTING=True)
        sent, info = send_smtp_email("student@example.com", "Test Subject", "<p>Hello</p>", "Hello")
        self.assertTrue(sent)
        self.assertEqual(info, "mock-smtp-dispatched")

        # 3. send_otp_code_email helper
        sent, info = send_otp_code_email("faculty@example.com", "123456", purpose="login")
        self.assertTrue(sent)
        self.assertEqual(info, "mock-smtp-dispatched")

        # 4. Mocking smtplib.SMTP for live transport simulation
        with patch('app.utils.email_service.get_smtp_config') as mock_cfg:
            mock_cfg.return_value = {
                'host': 'smtp.gmail.com',
                'port': 587,
                'user': 'pk0403564@gmail.com',
                'password': 'mock-app-password',
                'use_tls': True,
                'use_ssl': False,
                'from_email': 'SSVS Verification <pk0403564@gmail.com>',
                'is_testing': False
            }
            with patch('smtplib.SMTP') as mock_smtp_cls:
                mock_server = MagicMock()
                mock_smtp_cls.return_value.__enter__.return_value = mock_server

                sent, status = send_smtp_email("test@example.com", "Live Subject", "<b>Content</b>")
                self.assertTrue(sent)
                self.assertEqual(status, "smtp-dispatched")
                mock_server.starttls.assert_called_once()
                mock_server.login.assert_called_once_with('pk0403564@gmail.com', 'mock-app-password')
                mock_server.send_message.assert_called_once()

if __name__ == '__main__':
    unittest.main()

