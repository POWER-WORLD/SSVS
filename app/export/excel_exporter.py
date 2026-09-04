from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.models.form import Form
from app.models.submission import Submission
from app.models.scoring import LeaderboardEntry, CalculatedScore

class ExcelExporter:
    """Generates professional, multi-sheet formatted Excel reports."""

    # Color tokens
    HEADER_FILL = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")  # Deep Indigo
    SUBHEADER_FILL = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    SECTION_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    # Medal highlights
    GOLD_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    SILVER_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    BRONZE_FILL = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    
    # Fonts
    HEADER_FONT = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    COL_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    BOLD_FONT = Font(name="Calibri", size=11, bold=True)
    NORMAL_FONT = Font(name="Calibri", size=11)
    
    # Borders
    THIN_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    @classmethod
    def generate_form_report(cls, form: Form) -> BytesIO:
        wb = openpyxl.Workbook()
        
        # 1. Leaderboard Sheet
        ws_leaderboard = wb.active
        ws_leaderboard.title = "Leaderboard & Ranks"
        cls._build_leaderboard_sheet(ws_leaderboard, form)
        
        # 2. Raw Submissions Sheet
        ws_raw = wb.create_sheet(title="Raw Submissions")
        cls._build_raw_submissions_sheet(ws_raw, form)
        
        # 3. Platform Statistics Sheet
        ws_stats = wb.create_sheet(title="Coding Platform Stats")
        cls._build_platform_stats_sheet(ws_stats, form)
        
        # 4. Analytics Summary Sheet
        ws_analytics = wb.create_sheet(title="Executive Analytics")
        cls._build_analytics_sheet(ws_analytics, form)
        
        # Save to buffer
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @classmethod
    def _apply_auto_width(cls, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if cell.row <= 3 and len(val) > 30:
                    continue  # skip title banner
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def _build_leaderboard_sheet(cls, ws, form: Form):
        ws.views.sheetView[0].showGridLines = True
        
        # Header banner
        ws.merge_cells("A1:K1")
        top_cell = ws["A1"]
        top_cell.value = f"STUDENT SCORE VIEW SYSTEM — {form.title.upper()}"
        top_cell.font = cls.HEADER_FONT
        top_cell.fill = cls.HEADER_FILL
        top_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40
        
        # Metadata banner
        ws.merge_cells("A2:K2")
        sub_cell = ws["A2"]
        sub_cell.value = f"College: {form.teacher.college_name} | Department: {form.teacher.department} | Teacher: {form.teacher.full_name} | Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
        sub_cell.fill = cls.SUBHEADER_FILL
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 24
        
        # Table column titles
        headers = ["Rank", "Student Name", "Roll Number", "Department", "Semester", "CGPA", "Coding Score", "GitHub Score", "Total Score (100)", "Percentile", "Grade"]
        ws.append([]) # row 3 empty
        ws.append(headers) # row 4
        ws.row_dimensions[4].height = 26
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = cls.COL_HEADER_FONT
            cell.fill = cls.SUBHEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cls.THIN_BORDER

        # Entries
        entries = LeaderboardEntry.query.filter_by(form_id=form.id).order_by(LeaderboardEntry.rank.asc()).all()
        for i, entry in enumerate(entries):
            sub = entry.submission
            calc = sub.latest_score if sub else None
            grade = calc.grade if calc else "N/A"
            
            row = [
                f"#{entry.rank}",
                entry.student_name,
                entry.roll_number,
                entry.department or 'N/A',
                sub.semester if sub else 'N/A',
                f"{entry.cgpa:.2f}" if entry.cgpa else "0.00",
                f"{entry.coding_score:.1f}",
                f"{entry.github_score:.1f}",
                f"{entry.total_score:.2f}",
                f"{entry.percentile:.1f}%",
                grade
            ]
            ws.append(row)
            curr_row = ws.max_row
            ws.row_dimensions[curr_row].height = 22
            
            # Row styling
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = cls.NORMAL_FONT
                cell.border = cls.THIN_BORDER
                cell.alignment = Alignment(horizontal="center" if col_idx not in [2, 4] else "left", vertical="center")
                
                # Highlight podium
                if entry.rank == 1:
                    cell.fill = cls.GOLD_FILL
                elif entry.rank == 2:
                    cell.fill = cls.SILVER_FILL
                elif entry.rank == 3:
                    cell.fill = cls.BRONZE_FILL
                elif i % 2 == 1:
                    cell.fill = cls.ZEBRA_FILL

        ws.freeze_panes = "A5"
        cls._apply_auto_width(ws)

    @classmethod
    def _build_raw_submissions_sheet(cls, ws, form: Form):
        ws.views.sheetView[0].showGridLines = True
        
        headers = ["Submission ID", "Student Name", "Roll Number", "Email", "Phone", "Department", "Semester", "Batch", "CGPA", "Backlogs", "Status", "Submitted At"]
        ws.append(headers)
        ws.row_dimensions[1].height = 26
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = cls.COL_HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cls.THIN_BORDER

        submissions = form.submissions.order_by(Submission.submitted_at.desc()).all()
        for i, s in enumerate(submissions):
            row = [
                s.uuid[:8],
                s.student_name,
                s.roll_number,
                s.email,
                s.phone or '',
                s.department or '',
                s.semester or '',
                s.batch or '',
                s.cgpa or 0.0,
                s.backlogs or 0,
                s.status.capitalize(),
                s.submitted_at.strftime('%Y-%m-%d %H:%M') if s.submitted_at else ''
            ]
            ws.append(row)
            curr_row = ws.max_row
            ws.row_dimensions[curr_row].height = 20
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = cls.NORMAL_FONT
                cell.border = cls.THIN_BORDER
                cell.alignment = Alignment(horizontal="center" if col_idx not in [2, 4] else "left", vertical="center")
                if i % 2 == 1:
                    cell.fill = cls.ZEBRA_FILL

        ws.freeze_panes = "A2"
        cls._apply_auto_width(ws)

    @classmethod
    def generate_custom_export(cls, form: Form, submission_ids: list = None, column_keys: list = None) -> BytesIO:
        """
        Generate a custom Excel workbook containing ONLY the teacher-selected rows and columns.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Custom Export"
        ws.views.sheetView[0].showGridLines = True

        # Resolve submissions
        query = form.submissions
        if submission_ids and len(submission_ids) > 0:
            # Clean IDs
            clean_ids = []
            for sid in submission_ids:
                try:
                    clean_ids.append(int(sid))
                except (ValueError, TypeError):
                    pass
            if clean_ids:
                query = query.filter(Submission.id.in_(clean_ids))

        submissions = query.order_by(Submission.submitted_at.desc()).all()

        # Build column definitions dictionary
        col_defs = {
            'rank': {'title': 'Rank', 'align': 'center'},
            'student_name': {'title': 'Student Name', 'align': 'left'},
            'roll_number': {'title': 'Roll Number', 'align': 'center'},
            'email': {'title': 'Email Address', 'align': 'left'},
            'phone': {'title': 'Phone', 'align': 'center'},
            'department': {'title': 'Department', 'align': 'left'},
            'semester': {'title': 'Semester', 'align': 'center'},
            'batch': {'title': 'Batch', 'align': 'center'},
            'cgpa': {'title': 'CGPA', 'align': 'center'},
            'backlogs': {'title': 'Active Backlogs', 'align': 'center'},
            'total_score': {'title': 'Total Score (100)', 'align': 'center'},
            'percentile': {'title': 'Percentile', 'align': 'center'},
            'grade': {'title': 'Grade', 'align': 'center'},
            'verification_status': {'title': 'Status', 'align': 'center'},
            'teacher_remarks': {'title': 'Teacher Remarks', 'align': 'left'},
            'submitted_at': {'title': 'Submitted Date', 'align': 'center'},
            # LeetCode
            'leetcode_handle': {'title': 'LeetCode Handle', 'align': 'left'},
            'leetcode_solved': {'title': 'LeetCode Solved', 'align': 'center'},
            'leetcode_rating': {'title': 'LeetCode Rating', 'align': 'center'},
            'leetcode_rank': {'title': 'LeetCode Rank', 'align': 'center'},
            # CodeChef
            'codechef_handle': {'title': 'CodeChef Handle', 'align': 'left'},
            'codechef_rating': {'title': 'CodeChef Rating', 'align': 'center'},
            'codechef_stars': {'title': 'CodeChef Stars', 'align': 'center'},
            'codechef_solved': {'title': 'CodeChef Solved', 'align': 'center'},
            # Codeforces
            'codeforces_handle': {'title': 'Codeforces Handle', 'align': 'left'},
            'codeforces_rating': {'title': 'Codeforces Rating', 'align': 'center'},
            'codeforces_rank': {'title': 'Codeforces Rank', 'align': 'center'},
            # GitHub
            'github_handle': {'title': 'GitHub Handle', 'align': 'left'},
            'github_repos': {'title': 'GitHub Repos', 'align': 'center'},
            'github_stars': {'title': 'GitHub Stars', 'align': 'center'},
            'github_contributions': {'title': 'GitHub Contributions', 'align': 'center'},
            # HackerRank
            'hackerrank_handle': {'title': 'HackerRank Handle', 'align': 'left'},
            'hackerrank_badges': {'title': 'HackerRank Badges', 'align': 'center'},
            # GeeksforGeeks
            'gfg_handle': {'title': 'GFG Handle', 'align': 'left'},
            'gfg_score': {'title': 'GFG Coding Score', 'align': 'center'},
            'gfg_problems': {'title': 'GFG Solved', 'align': 'center'},
            # AtCoder
            'atcoder_handle': {'title': 'AtCoder Handle', 'align': 'left'},
            'atcoder_rating': {'title': 'AtCoder Rating', 'align': 'center'},
            'atcoder_rank': {'title': 'AtCoder Rank', 'align': 'center'},
            # InterviewBit
            'interviewbit_handle': {'title': 'InterviewBit Handle', 'align': 'left'},
            'interviewbit_score': {'title': 'InterviewBit Score', 'align': 'center'},
            'interviewbit_solved': {'title': 'InterviewBit Solved', 'align': 'center'},
            # Kaggle
            'kaggle_handle': {'title': 'Kaggle Handle', 'align': 'left'},
            'kaggle_tier': {'title': 'Kaggle Tier', 'align': 'center'},
            'kaggle_medals': {'title': 'Kaggle Medals', 'align': 'center'},
        }

        # Add dynamic custom form fields to available columns
        for f in form.fields.order_by('display_order').all():
            k = f"q_{f.field_key}"
            if k not in col_defs:
                col_defs[k] = {'title': f.label, 'align': 'left', 'field_key': f.field_key}

        # Determine chosen columns
        if not column_keys or len(column_keys) == 0:
            active_cols = ['rank', 'student_name', 'roll_number', 'department', 'cgpa', 'total_score', 'percentile', 'grade']
        else:
            active_cols = [c for c in column_keys if c in col_defs]
            if not active_cols:
                active_cols = ['rank', 'student_name', 'roll_number', 'cgpa', 'total_score', 'grade']

        num_cols = len(active_cols)
        last_col_letter = get_column_letter(num_cols)

        # Header banner
        ws.merge_cells(f"A1:{last_col_letter}1")
        top_cell = ws["A1"]
        top_cell.value = f"CUSTOM COHORT EXPORT — {form.title.upper()}"
        top_cell.font = cls.HEADER_FONT
        top_cell.fill = cls.HEADER_FILL
        top_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # Subtitle
        ws.merge_cells(f"A2:{last_col_letter}2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Exported: {len(submissions)} Students | {num_cols} Columns | Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
        sub_cell.fill = cls.SUBHEADER_FILL
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        # Column headers row (Row 3)
        headers = [col_defs[c]['title'] for c in active_cols]
        ws.append(headers)
        ws.row_dimensions[3].height = 26

        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.font = cls.COL_HEADER_FONT
            cell.fill = cls.SUBHEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cls.THIN_BORDER

        # Rows
        for i, s in enumerate(submissions):
            score = s.latest_score
            profiles = {p.platform_name.lower(): p for p in s.platform_profiles.all()}
            
            row_data = []
            for c in active_cols:
                val = ""
                # Basic
                if c == 'rank':
                    val = f"#{score.rank}" if score and score.rank else "—"
                elif c == 'student_name':
                    val = s.student_name
                elif c == 'roll_number':
                    val = s.roll_number
                elif c == 'email':
                    val = s.email
                elif c == 'phone':
                    val = s.phone or ''
                elif c == 'department':
                    val = s.department or ''
                elif c == 'semester':
                    val = s.semester or ''
                elif c == 'batch':
                    val = s.batch or ''
                elif c == 'cgpa':
                    val = f"{s.cgpa:.2f}" if s.cgpa else "0.00"
                elif c == 'backlogs':
                    val = s.backlogs or 0
                elif c == 'total_score':
                    val = f"{score.total_score:.2f}" if score else "0.00"
                elif c == 'percentile':
                    val = f"{score.percentile:.1f}%" if score else "0.0%"
                elif c == 'grade':
                    val = score.grade if score else "—"
                elif c == 'verification_status':
                    val = (getattr(s, 'verification_status', None) or s.status).capitalize()
                elif c == 'teacher_remarks':
                    val = getattr(s, 'teacher_remarks', '') or ''
                elif c == 'submitted_at':
                    val = s.submitted_at.strftime('%Y-%m-%d %H:%M') if s.submitted_at else ''
                
                # LeetCode
                elif c == 'leetcode_handle':
                    val = profiles.get('leetcode').username if 'leetcode' in profiles else ''
                elif c == 'leetcode_solved':
                    val = profiles.get('leetcode').problems_solved if 'leetcode' in profiles else 0
                elif c == 'leetcode_rating':
                    val = profiles.get('leetcode').contest_rating if 'leetcode' in profiles else 0
                elif c == 'leetcode_rank':
                    val = profiles.get('leetcode').global_rank if 'leetcode' in profiles else 0
                
                # CodeChef
                elif c == 'codechef_handle':
                    val = profiles.get('codechef').username if 'codechef' in profiles else ''
                elif c == 'codechef_rating':
                    val = int(profiles.get('codechef').contest_rating) if 'codechef' in profiles else 0
                elif c == 'codechef_stars':
                    val = profiles.get('codechef').stars_or_badges if 'codechef' in profiles else 0
                elif c == 'codechef_solved':
                    val = profiles.get('codechef').problems_solved if 'codechef' in profiles else 0

                # Codeforces
                elif c == 'codeforces_handle':
                    val = profiles.get('codeforces').username if 'codeforces' in profiles else ''
                elif c == 'codeforces_rating':
                    val = int(profiles.get('codeforces').contest_rating) if 'codeforces' in profiles else 0
                elif c == 'codeforces_rank':
                    val = (profiles.get('codeforces').raw_data or {}).get('rank', 'unrated') if 'codeforces' in profiles else ''

                # GitHub
                elif c == 'github_handle':
                    val = profiles.get('github').username if 'github' in profiles else ''
                elif c == 'github_repos':
                    val = profiles.get('github').problems_solved if 'github' in profiles else 0
                elif c == 'github_stars':
                    val = (profiles.get('github').raw_data or {}).get('stars_earned', 0) if 'github' in profiles else 0
                elif c == 'github_contributions':
                    val = (profiles.get('github').raw_data or {}).get('contributions_year', 0) if 'github' in profiles else 0

                # HackerRank
                elif c == 'hackerrank_handle':
                    val = profiles.get('hackerrank').username if 'hackerrank' in profiles else ''
                elif c == 'hackerrank_badges':
                    val = profiles.get('hackerrank').stars_or_badges if 'hackerrank' in profiles else 0

                # GeeksforGeeks
                elif c == 'gfg_handle':
                    p = profiles.get('gfg') or profiles.get('geeksforgeeks')
                    val = p.username if p else ''
                elif c == 'gfg_score':
                    p = profiles.get('gfg') or profiles.get('geeksforgeeks')
                    val = (p.raw_data or {}).get('coding_score', 0) if p else 0
                elif c == 'gfg_problems':
                    p = profiles.get('gfg') or profiles.get('geeksforgeeks')
                    val = p.problems_solved if p else 0

                # AtCoder
                elif c == 'atcoder_handle':
                    val = profiles.get('atcoder').username if 'atcoder' in profiles else ''
                elif c == 'atcoder_rating':
                    val = int(profiles.get('atcoder').contest_rating) if 'atcoder' in profiles else 0
                elif c == 'atcoder_rank':
                    val = profiles.get('atcoder').global_rank if 'atcoder' in profiles else 0

                # InterviewBit
                elif c == 'interviewbit_handle':
                    val = profiles.get('interviewbit').username if 'interviewbit' in profiles else ''
                elif c == 'interviewbit_score':
                    val = int(profiles.get('interviewbit').contest_rating) if 'interviewbit' in profiles else 0
                elif c == 'interviewbit_solved':
                    val = profiles.get('interviewbit').problems_solved if 'interviewbit' in profiles else 0

                # Kaggle
                elif c == 'kaggle_handle':
                    val = profiles.get('kaggle').username if 'kaggle' in profiles else ''
                elif c == 'kaggle_tier':
                    val = (profiles.get('kaggle').raw_data or {}).get('tier', 'Contributor') if 'kaggle' in profiles else ''
                elif c == 'kaggle_medals':
                    val = profiles.get('kaggle').stars_or_badges if 'kaggle' in profiles else 0

                # Dynamic questions
                elif c.startswith('q_'):
                    f_key = col_defs[c].get('field_key')
                    val = s.get_value(f_key) or ''
                
                row_data.append(val)

            ws.append(row_data)
            curr_row = ws.max_row
            ws.row_dimensions[curr_row].height = 20

            for col_idx in range(1, num_cols + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = cls.NORMAL_FONT
                cell.border = cls.THIN_BORDER
                cell.alignment = Alignment(horizontal=col_defs[active_cols[col_idx - 1]]['align'], vertical="center")
                if i % 2 == 1:
                    cell.fill = cls.ZEBRA_FILL

        ws.freeze_panes = "A4"
        cls._apply_auto_width(ws)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @classmethod
    def _build_platform_stats_sheet(cls, ws, form: Form):
        ws.views.sheetView[0].showGridLines = True
        
        headers = [
            "Roll No", "Student Name",
            "LeetCode Solved", "LeetCode Rating",
            "CodeChef Rating", "CodeChef Stars",
            "Codeforces Rating", "Codeforces Rank",
            "GitHub Repos", "GitHub Stars",
            "HackerRank Badges",
            "GFG Score", "AtCoder Rating", "InterviewBit Score", "Kaggle Medals"
        ]
        ws.append(headers)
        ws.row_dimensions[1].height = 26
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = cls.COL_HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cls.THIN_BORDER

        for i, s in enumerate(form.submissions.all()):
            lc_solved = lc_rating = cc_rating = cc_stars = cf_rating = gh_repos = gh_stars = hr_badges = 0
            gfg_score = atc_rating = ib_score = kaggle_medals = 0
            cf_rank = 'unrated'
            
            for p in s.platform_profiles.all():
                raw = p.raw_data or {}
                p_name = p.platform_name.lower()
                if p_name == 'leetcode':
                    lc_solved = p.problems_solved or raw.get('total_solved', 0)
                    lc_rating = p.contest_rating or raw.get('contest_rating', 0)
                elif p_name == 'codechef':
                    cc_rating = p.contest_rating or raw.get('current_rating', 0)
                    cc_stars = p.stars_or_badges or raw.get('stars', 0)
                elif p_name == 'codeforces':
                    cf_rating = p.contest_rating or raw.get('rating', 0)
                    cf_rank = raw.get('rank', 'unrated')
                elif p_name == 'github':
                    gh_repos = raw.get('public_repos', 0)
                    gh_stars = raw.get('stars_earned', 0)
                elif p_name == 'hackerrank':
                    hr_badges = raw.get('badges_count', 0)
                elif p_name in ('gfg', 'geeksforgeeks'):
                    gfg_score = raw.get('coding_score', 0) or p.problems_solved
                elif p_name == 'atcoder':
                    atc_rating = p.contest_rating or raw.get('current_rating', 0)
                elif p_name == 'interviewbit':
                    ib_score = raw.get('score', 0) or p.contest_rating
                elif p_name == 'kaggle':
                    kaggle_medals = p.stars_or_badges or raw.get('total_medals', 0)

            row = [
                s.roll_number, s.student_name,
                lc_solved, f"{lc_rating:.1f}" if lc_rating else "0",
                int(cc_rating), int(cc_stars),
                int(cf_rating), cf_rank,
                int(gh_repos), int(gh_stars),
                int(hr_badges),
                int(gfg_score), int(atc_rating), int(ib_score), int(kaggle_medals)
            ]
            ws.append(row)
            curr_row = ws.max_row
            ws.row_dimensions[curr_row].height = 20
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = cls.NORMAL_FONT
                cell.border = cls.THIN_BORDER
                cell.alignment = Alignment(horizontal="center" if col_idx not in [2] else "left", vertical="center")
                if i % 2 == 1:
                    cell.fill = cls.ZEBRA_FILL

        ws.freeze_panes = "A2"
        cls._apply_auto_width(ws)

    @classmethod
    def _build_analytics_sheet(cls, ws, form: Form):
        ws.views.sheetView[0].showGridLines = True
        
        ws.merge_cells("A1:D1")
        title_cell = ws["A1"]
        title_cell.value = "EXECUTIVE COHORT PERFORMANCE SUMMARY"
        title_cell.font = cls.HEADER_FONT
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36
        
        entries = LeaderboardEntry.query.filter_by(form_id=form.id).all()
        total_count = len(entries)
        scores = [e.total_score for e in entries] if entries else []
        cgpas = [e.cgpa for e in entries if e.cgpa] if entries else []
        
        avg_score = sum(scores) / total_count if total_count else 0.0
        high_score = max(scores) if scores else 0.0
        low_score = min(scores) if scores else 0.0
        median_score = sorted(scores)[total_count // 2] if total_count else 0.0
        avg_cgpa = sum(cgpas) / len(cgpas) if cgpas else 0.0
        
        ws.append([])
        metrics = [
            ("Total Submissions Evaluated", total_count),
            ("Cohort Highest Score", f"{high_score:.2f} / 100"),
            ("Cohort Lowest Score", f"{low_score:.2f} / 100"),
            ("Cohort Mean Score", f"{avg_score:.2f} / 100"),
            ("Cohort Median Score", f"{median_score:.2f} / 100"),
            ("Average Student CGPA", f"{avg_cgpa:.2f} / 10.0"),
            ("Active Leaderboard Submissions", total_count)
        ]
        
        ws.append(["Performance Indicator", "Metric Value"])
        ws.row_dimensions[3].height = 24
        ws.cell(row=3, column=1).font = cls.COL_HEADER_FONT
        ws.cell(row=3, column=1).fill = cls.SUBHEADER_FILL
        ws.cell(row=3, column=2).font = cls.COL_HEADER_FONT
        ws.cell(row=3, column=2).fill = cls.SUBHEADER_FILL
        
        for label, val in metrics:
            ws.append([label, str(val)])
            r = ws.max_row
            ws.row_dimensions[r].height = 20
            c1 = ws.cell(row=r, column=1)
            c2 = ws.cell(row=r, column=2)
            c1.font = cls.BOLD_FONT
            c2.font = cls.NORMAL_FONT
            c1.border = cls.THIN_BORDER
            c2.border = cls.THIN_BORDER
            c1.fill = cls.SECTION_FILL
            c2.alignment = Alignment(horizontal="center")
            
        cls._apply_auto_width(ws)
